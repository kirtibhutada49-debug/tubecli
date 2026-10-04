# -*- coding: utf-8 -*-
"""write_doc (nút Lưu của node Word trên Flow) phải SỬA TẠI CHỖ, không dựng lại tài liệu.

4/10/2026: write_doc cũ tạo `Document()` mới tinh rồi đắp đoạn vào → mở rồi bấm Lưu là mất bảng, ảnh, header/footer,
lề trang, style riêng và định dạng trộn trong đoạn. Test dựng một file .docx thật có đủ những thứ đó, giả lập đúng
luồng của trình sửa (read_doc → sửa vài đoạn → write_doc) rồi mở lại kiểm từng thứ.

Chạy:  python tests/write_doc_inplace_test.py   (exit 0 = pass)
"""
import base64
import os
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except AttributeError:
    pass

from docx import Document                                        # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                    # noqa: E402
from docx.oxml import OxmlElement                                # noqa: E402
from docx.oxml.ns import qn                                      # noqa: E402
from docx.shared import Cm, Pt                                   # noqa: E402

from tubecli.extensions.file_manager.file_service import FileService  # noqa: E402

PASS = FAIL = 0


def check(name, ok, detail=""):
    global PASS, FAIL
    if ok:
        PASS += 1
        print(f"[PASS] {name}")
    else:
        FAIL += 1
        print(f"[FAIL] {name} -> {detail}")


PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def add_hyperlink(par, url, text):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


def build(path, img):
    d = Document()
    sec = d.sections[0]
    sec.top_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(3), Cm(1.5)
    sec.header.paragraphs[0].text = "ĐẦU TRANG CÔNG TY"
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(13)
    p = d.add_paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")               # 0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p = d.add_paragraph()                                                   # 1: «1.1. » đậm + thân thường
    r = p.add_run("1.1. ")
    r.bold = True
    p.add_run("Bên A cam kết thanh toán đúng hạn.")
    p = d.add_paragraph("Điều 1.\tNội dung điều một")                       # 2: có tab
    p = d.add_paragraph("Xem chi tiết tại ")                               # 3: có liên kết
    add_hyperlink(p, "https://tubecli.app", "tubecli.app")
    p.add_run(" để biết thêm.")
    t = d.add_table(rows=2, cols=2)                                         # bảng
    t.cell(0, 0).text, t.cell(0, 1).text = "Hạng mục", "Đơn giá"
    t.cell(1, 0).text, t.cell(1, 1).text = "Thiết kế", "5.000.000"
    d.add_paragraph().add_run().add_picture(img, width=Cm(1))               # 4: ảnh
    d.add_paragraph("Đoạn cuối giữ nguyên.")                                # 5
    d.save(path)


def runs_of(par):
    return [(r.text, bool(r.bold)) for r in FileService._para_runs(par)]


def main():
    # Thư mục tạm TRONG repo: %TEMP% (~/AppData/Local) nằm trong BLOCKED_PATHS của File Manager.
    tmp = Path(tempfile.mkdtemp(prefix="_write_doc_test_", dir=str(ROOT / "tests")))
    try:
        img = tmp / "dot.png"
        img.write_bytes(PNG)
        f = tmp / "hop-dong.docx"
        build(str(f), str(img))
        svc = FileService(extra_roots=[str(tmp)])

        # ── 1. Mở rồi Lưu, KHÔNG sửa gì → tài liệu không đổi gì ───────────────
        data = svc.read_doc(str(f))
        before = Document(str(f))
        res = svc.write_doc(str(f), data["paragraphs"])
        after = Document(str(f))
        check("1a Lưu không sửa: báo sửa tại chỗ, 0 đoạn đổi", res.get("in_place") and res["changed"] == 0 and res["added"] == 0, res)
        check("1b bảng còn nguyên (1 bảng, 4 ô đúng chữ)", len(after.tables) == 1
              and [c.text for row in after.tables[0].rows for c in row.cells] == ["Hạng mục", "Đơn giá", "Thiết kế", "5.000.000"],
              len(after.tables))
        check("1c ảnh còn nguyên", len(after.inline_shapes) == 1, len(after.inline_shapes))
        check("1d header còn nguyên", after.sections[0].header.paragraphs[0].text == "ĐẦU TRANG CÔNG TY")
        check("1e lề trang còn nguyên (trái 3 cm, phải 1,5 cm)", abs(after.sections[0].left_margin.cm - 3) < 0.01
              and abs(after.sections[0].right_margin.cm - 1.5) < 0.01)
        check("1f định dạng trộn «1.1.» đậm + thân thường còn nguyên", runs_of(after.paragraphs[1]) == runs_of(before.paragraphs[1])
              and runs_of(after.paragraphs[1])[0] == ("1.1. ", True) and runs_of(after.paragraphs[1])[1][1] is False,
              runs_of(after.paragraphs[1]))
        check("1g liên kết còn nguyên", "tubecli.app" in after.paragraphs[3].text and after.paragraphs[3]._p.find(qn("w:hyperlink")) is not None)

        # ── 2. Sửa như người dùng: đổi chữ giữa đoạn, đổi căn lề, thêm đoạn ───
        paras = svc.read_doc(str(f))["paragraphs"]
        paras[1]["text"] = "1.1. Bên A cam kết thanh toán trong 5 ngày."         # đổi khúc cuối của phần thường
        paras[2]["text"] = "Điều 1.\tNội dung điều một đã sửa"                   # đoạn có tab
        paras[3]["text"] = "Xem chi tiết tại tubecli.app để biết thêm nhé."     # đuôi sau liên kết
        paras[5]["align"] = "right"
        paras.append({"text": "Đoạn mới thêm.", "style": "Normal", "align": "justify", "bold": False, "italic": False,
                      "underline": False, "size": 13})
        res = svc.write_doc(str(f), paras)
        d2 = Document(str(f))
        check("2a đếm đúng: 4 đoạn đổi, 1 đoạn thêm, không gộp run", res["changed"] == 4 and res["added"] == 1
              and res["runs_merged"] == 0, res)
        check("2b chữ «1.1. » vẫn đậm, phần sửa nằm trong run thường", runs_of(d2.paragraphs[1])
              == [("1.1. ", True), ("Bên A cam kết thanh toán trong 5 ngày.", False)], runs_of(d2.paragraphs[1]))
        check("2c đoạn có tab: không nhân đôi chữ (bẫy OfficeCLI)", d2.paragraphs[2].text == "Điều 1.\tNội dung điều một đã sửa",
              repr(d2.paragraphs[2].text))
        check("2d sửa đuôi sau liên kết: liên kết còn, chữ đúng", d2.paragraphs[3].text == "Xem chi tiết tại tubecli.app để biết thêm nhé."
              and d2.paragraphs[3]._p.find(qn("w:hyperlink")) is not None, d2.paragraphs[3].text)
        check("2e căn phải áp đúng đoạn", d2.paragraphs[5].alignment == WD_ALIGN_PARAGRAPH.RIGHT)
        check("2f đoạn mới ở CUỐI, sau đoạn cũ, bảng/ảnh/header/lề vẫn còn",
              d2.paragraphs[-1].text == "Đoạn mới thêm." and d2.paragraphs[-2].text == "Đoạn cuối giữ nguyên."
              and len(d2.tables) == 1 and len(d2.inline_shapes) == 1
              and d2.sections[0].header.paragraphs[0].text == "ĐẦU TRANG CÔNG TY" and abs(d2.sections[0].left_margin.cm - 3) < 0.01)
        check("2g đoạn không sửa giữ nguyên chữ + căn giữa + đậm", d2.paragraphs[0].text == "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
              and d2.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER and d2.paragraphs[0].runs[0].bold)

        # ── 3. Đổi định dạng cả đoạn ─────────────────────────────────────────
        paras = svc.read_doc(str(f))["paragraphs"]
        paras[5]["bold"] = True
        paras[5]["size"] = 14
        res = svc.write_doc(str(f), paras)
        d3 = Document(str(f))
        check("3a đậm + cỡ 14 áp vào đúng đoạn", all(r.bold for r in d3.paragraphs[5].runs if r.text)
              and all(r.font.size and abs(r.font.size.pt - 14) < 0.05 for r in d3.paragraphs[5].runs if r.text)
              and res["changed"] == 1, res)

        # ── 4. Client gửi THIẾU đoạn → không xoá gì ─────────────────────────
        n = len(Document(str(f)).paragraphs)
        svc.write_doc(str(f), svc.read_doc(str(f))["paragraphs"][:2])
        check("4 gửi thiếu đoạn: số đoạn giữ nguyên, không xoá", len(Document(str(f)).paragraphs) == n)

        # ── 5. Chèn thuần tuý / xoá thuần tuý / xoá hết chữ ─────────────────
        paras = svc.read_doc(str(f))["paragraphs"]
        paras[1]["text"] = "1.1. 2. Bên A cam kết thanh toán trong 5 ngày."     # chèn ngay sau «1.1. » (đậm)
        paras[0]["text"] = "CỘNG HÒA VIỆT NAM"                                    # xoá khúc giữa
        paras[6]["text"] = ""                                                    # xoá hết chữ đoạn mới
        svc.write_doc(str(f), paras)
        d5 = Document(str(f))
        check("5a chèn ngay sau phần đậm: ăn theo ký tự trước (đậm), phần thường không đổi",
              runs_of(d5.paragraphs[1]) == [("1.1. 2. ", True), ("Bên A cam kết thanh toán trong 5 ngày.", False)], runs_of(d5.paragraphs[1]))
        check("5b xoá khúc giữa", d5.paragraphs[0].text == "CỘNG HÒA VIỆT NAM" and d5.paragraphs[0].runs[0].bold)
        check("5c xoá hết chữ một đoạn: đoạn rỗng nhưng vẫn còn đoạn", d5.paragraphs[6].text == "" and len(d5.paragraphs) == n)

        # ── 6. File chưa có → tạo mới ────────────────────────────────────────
        nf = tmp / "moi.docx"
        res = svc.write_doc(str(nf), [{"text": "Xin chào", "style": "Normal", "align": "center", "bold": True}])
        dn = Document(str(nf))
        check("6 file chưa tồn tại → tạo mới, đúng chữ + căn giữa + đậm", not res["in_place"] and dn.paragraphs[0].text == "Xin chào"
              and dn.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER and dn.paragraphs[0].runs[0].bold)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
    print(f"\n{PASS} pass, {FAIL} fail")
    sys.exit(1 if FAIL else 0)
